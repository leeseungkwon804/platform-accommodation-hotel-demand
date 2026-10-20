"""Build the manuscript on top of the MDPI Sustainability template (converted .dot -> .docx base).
Fills only sections whose source exists; unknown front/back-matter fields are left empty (no placeholder text).
Citations [@key] are numbered by first appearance. usage: python -I 20_build_manuscript.py PROJECT_DIR"""
import sys, re, json, copy, pathlib, zipfile, importlib.util
import pandas as pd
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

R = pathlib.Path(sys.argv[1])
BASE = R / "manuscript/_template/sustainability-template_base.docx"
OUT = R / "manuscript/C10_manuscript_draft.docx"
F = json.load(open(R / "out/facts_methods.json"))
ev_path = R / "data/derived/regulation_events.csv"
EV = None
if ev_path.exists():
    ev = pd.read_csv(ev_path)
    EV = int((ev.source_level.astype(str).str.strip() == "1차").sum()) or None
spec = importlib.util.spec_from_file_location("methods", R / "manuscript/src/methods.py")
M = importlib.util.module_from_spec(spec); spec.loader.exec_module(M)
REF_DOI = json.load(open(R / "manuscript/src/refs_verified.json"))
REF_X = json.load(open(R / "manuscript/src/refs_extra.json"))
spec = importlib.util.spec_from_file_location("text", R / "manuscript/src/text.py")
T = importlib.util.module_from_spec(spec); spec.loader.exec_module(T)
O = R / "out"
def jl(p): return json.load(open(p))
X = dict(F=F, evd=jl(O / "L1_event_dynamics.json"), c2=jl(O / "C2_claim1_defence.json"),
         th={r["tag"].split()[0]: r for r in jl(O / "L1_threshold.json")}, tf={r["tag"].split()[0]: r for r in jl(O / "L1_threshold_fixed.json")},
         cf_conv=jl(O / "L1_causal_forest.json"), cf_conv_pl=jl(O / "L1_causal_forest_placebo.json"),
         cf_fixed=jl(O / "L1_causal_forest_pen_fixed.json"), cf_fixed_pl=jl(O / "L1_causal_forest_pen_fixed_placebo.json"),
         l2=pd.read_csv(O / "L2_forecast_summary.csv"), l2lag=pd.read_csv(O / "L2_forecast_summary_lag3.csv"),
         l3pl=jl(O / "C1_L3_placebo.json"), covid=jl(O / "covid_share.json"),
         lev=jl(O / "L1_level_passthrough.json"), norm=jl(O / "L1_normlevel.json"), normgate=jl(O / "L1_normlevel_gate.json"),
         scale=jl(O / "scale_facts.json"), fig=jl(R / "figs/figure_facts.json")["facts"])
FIGDIR = R / "figs"

d = docx.Document(BASE)
body = d.element.body
paras = d.paragraphs

def set_text(p, text):
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    if p.runs: p.runs[0].text = text
    else: p.add_run(text)

# --- front matter: article type kept, everything unknown emptied ---
set_text(paras[0], "Article")
set_text(paras[1], "Platform Accommodation and Hotel Demand in European Regions with Evidence from Regulation and Placebo Tests")  # c10dec5: title 3, provisional
set_text(paras[2], "")                       # authors: not received
for p in paras[3:6]:                         # affiliations / correspondence: not received
    set_text(p, "")
for p in [paras[4], paras[5]]:
    p._p.getparent().remove(p._p)
abs_p = paras[7]; set_text(abs_p, T.abstract(X)[1])
# side text box (Academic Editor / Received / ... ) is filled by the editorial office; removed as in exemplars 01/02
for tbl in [t for t in body.iter(qn("w:tbl"))
            if "Academic Editor" in "".join(x.text or "" for x in t.iter(qn("w:t")))]:
    tbl.getparent().remove(tbl)  # floating side table (Academic Editor / dates / citation / copyright)
kw = paras[8]
kw_runs = kw.runs
if kw_runs:
    kw_runs[0].text = "Keywords: "
    for r in kw_runs[1:]: r._r.getparent().remove(r._r)
    kw.add_run(T.KEYWORDS[len("Keywords: "):])

# --- remove template body from "0. How to Use" through Appendix B text; keep References heading + notes ---
start = paras[10]._p; ref_head = paras[77]._p
el = start
while el is not None and el is not ref_head:
    nxt = el.getnext(); body.remove(el); el = nxt
# remove template reference examples (paragraphs after References heading, style MDPI81references)
for p in list(d.paragraphs):
    if p.style.style_id == "MDPI81references":
        p._p.getparent().remove(p._p)

# --- citation numbering ---
order = []
def cite(text):
    def rep(m):
        keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
        nums = []
        for k in keys:
            if k not in REF_DOI and k not in REF_X: raise KeyError(f"unverified reference key {k}")
            if k not in order: order.append(k)
            nums.append(order.index(k) + 1)
        nums = sorted(set(nums)); parts = []; i = 0
        while i < len(nums):
            j = i
            while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1: j += 1
            parts.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ",".join(map(str, nums[i:j + 1])))
            i = j + 1
        return "[" + ",".join(parts) + "]"
    return re.sub(r"\[(@[^\]]+)\]", rep, text)

def new_par(style, text, bold_label=None):
    p = OxmlElement("w:p"); pPr = OxmlElement("w:pPr"); ps = OxmlElement("w:pStyle"); ps.set(qn("w:val"), style)
    pPr.append(ps); p.append(pPr)
    def run(t, bold=False):
        r = OxmlElement("w:r")
        if bold:
            rPr = OxmlElement("w:rPr"); b = OxmlElement("w:b"); rPr.append(b); r.append(rPr)
        tt = OxmlElement("w:t"); tt.set(qn("xml:space"), "preserve"); tt.text = t; r.append(tt); return r
    if bold_label and text.startswith(bold_label):
        p.append(run(bold_label, True)); p.append(run(text[len(bold_label):]))
    elif text:
        p.append(run(text))
    return p

def new_table(rows):
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    t.style = d.styles["MDPI_4.1_three_line_table"]
    for i, row in enumerate(rows):
        tr = t.rows[i]._tr; trPr = tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit"); trPr.append(cs)
        if i == 0:
            th = OxmlElement("w:tblHeader"); trPr.append(th)
        for j, val in enumerate(row):
            c = t.cell(i, j); c.text = ""
            par = c.paragraphs[0]; par.style = d.styles["MDPI_4.2_table_body"]
            r = par.add_run(cite(val)); r.bold = (i == 0) or None
    tbl = t._tbl; tbl.getparent().remove(tbl); return tbl

def new_figure(name, caption):
    p = d.add_paragraph(); p.style = d.styles["MDPI_5.2_figure"]
    p.add_run().add_picture(str(FIGDIR / f"{name}.png"), width=docx.shared.Mm(180))
    el = p._p; el.getparent().remove(el)
    lab = caption.split(" ", 2); lab = f"{lab[0]} {lab[1]}"
    return [el, new_par("MDPI51figurecaption", cite(caption), bold_label=lab)]

FIG_ORDER = []
ITEMS = []
H = lambda style, text: ITEMS.append(new_par(style, text))
def emit(blocks):
    for style, text in blocks:
        if style == "TABLE": ITEMS.append(new_table(text))
        elif style == "FIGURE":
            FIG_ORDER.append(text[0]); ITEMS.extend(new_figure(*text))
        elif style in ("MDPI41tablecaption",):
            ITEMS.append(new_par(style, cite(text), bold_label=" ".join(text.split(" ", 2)[:2])))
        else: ITEMS.append(new_par(style, cite(text)))
H("MDPI21heading1", "1. Introduction")
emit(T.intro(X))
emit(M.blocks(F, EV, X))
H("MDPI21heading1", "3. Results"); emit(T.results(X))
H("MDPI21heading1", "4. Discussion"); emit(T.discussion(X))
H("MDPI21heading1", "5. Conclusions"); emit(T.conclusions(X))
for lab, text in [("Supplementary Materials:", T.supplementary(X)), ("Author Contributions:", "Author Contributions:"), ("Funding:", "Funding:"),
                  ("Institutional Review Board Statement:", M.IRB), ("Informed Consent Statement:", M.CONSENT),
                  ("Data Availability Statement:", M.DATA_AVAILABILITY), ("Conflicts of Interest:", "Conflicts of Interest:")]:
    ITEMS.append(new_par("MDPI62backmatter", text, bold_label=lab))
for it in ITEMS:
    ref_head.addprevious(it)

# --- references in citation order (style supplies numbering; no "1. " in text) ---
ISO4 = {"Tourism Management": "Tour. Manag.", "American Economic Review": "Am. Econ. Rev.", "Marketing Science": "Mark. Sci.",
        "Current Issues in Tourism": "Curr. Issues Tour.", "The Quarterly Journal of Economics": "Q. J. Econ.",
        "Review of Economics and Statistics": "Rev. Econ. Stat.", "Journal of Applied Econometrics": "J. Appl. Econom.",
        "Journal of Econometrics": "J. Econom.", "American Economic Review: Insights": "Am. Econ. Rev. Insights",
        "Journal of Marketing Research": "J. Mark. Res.", "Empirical Economics": "Empir. Econ.", "The Econometrics Journal": "Econom. J.",
        "Journal of the American Statistical Association": "J. Am. Stat. Assoc.", "The Annals of Statistics": "Ann. Stat.",
        "Journal of Business & Economic Statistics": "J. Bus. Econ. Stat.", "International Journal of Forecasting": "Int. J. Forecast.",
        "Nature Machine Intelligence": "Nat. Mach. Intell.", "Sustainability": "Sustainability"}

def ref_year(r):  # volume year: print year if registered; Sustainability volume v = year 2008 + v
    if r["journal"] == "Sustainability" and r.get("volume"): return 2008 + int(r["volume"])
    return r.get("year_print") or r["year"]

def fmt_doi(r):
    au = []
    for fam, giv in r["authors"]:
        toks = [t for t in re.split(r"\s+", giv or "") if t]
        part = []
        while toks and toks[-1][0].islower():  # nobiliary particles ("de la") belong to the family name
            part.insert(0, toks.pop())
        fam = " ".join(part + [fam])
        ini = "".join(".-".join(y[0] for y in t.split("-") if y) + "." for t in toks)
        au.append(f"{fam}, {ini}" if ini else fam)
    au_s = "; ".join(au)
    jn = r["journal"].replace("&amp;", "&"); j = ISO4.get(jn) or (r.get("short") or jn)
    vol = f", {r['volume']}" if r.get("volume") else ""
    page = (r.get("page") or "").replace("-", "–")
    pg = f", {page}" if page else (f", {r['article_number']}" if r.get("article_number") else "")
    title = r["title"].rstrip(".").rstrip("*").replace("&amp;", "&")
    tsep = " " if title.endswith("?") else ". "
    if r.get("type") == "proceedings-article":
        return f"{au_s} {title}{tsep}In {r['journal']}; {ref_year(r)}; pp. {page}. https://doi.org/{r['doi']}."
    return f"{au_s} {title}{tsep}{j} {ref_year(r)}{vol}{pg}. https://doi.org/{r['doi']}."
anchor = ref_head
for k in order:
    txt = fmt_doi(REF_DOI[k]) if k in REF_DOI else REF_X[k]
    p = new_par("MDPI81references", txt); anchor.addnext(p); anchor = p

d.save(OUT)

# --- post-build checks ---
z = zipfile.ZipFile(OUT); names = z.namelist()
rels = z.read("word/_rels/header3.xml.rels").decode()
chk = dict(
    header3_rId1_image3=bool(re.search(r'Id="rId1"[^>]*Target="media/image3.png"', rels)),
    header3_rId2_image4=bool(re.search(r'Id="rId2"[^>]*Target="media/image4.png"', rels)),
    image3_present="word/media/image3.png" in names, image4_present="word/media/image4.png" in names,
    comments=sum(1 for n in names if n == "word/comments.xml"),
    references=len(order))
xml = z.read("word/document.xml").decode()
txt = re.sub(r"<[^>]+>", "", xml)
chk["placeholder_hits"] = re.findall(r"TO BE|TBD|\[TO|XXX|Firstname|e-mail@e-mail", txt)
chk["ref_text_numbered"] = bool(re.search(r"MDPI81references.{0,400}?<w:t[^>]*>\s*\d+\.\s", xml))
chk["dup_accessed"] = len(re.findall(r"\(accessed on [^)]*\)\s*\(accessed on", txt))
chk["unresolved_cites"] = re.findall(r"\[@[^\]]+\]", txt)
print(json.dumps(chk, ensure_ascii=False, indent=1))
print("citation order:", order)

# --- figure/table numbering = first citation order; abstract length ---
body_txt = "\n".join(p.text for p in docx.Document(OUT).paragraphs)
first = {}
for m in re.finditer(r"(Figure|Table) (\d+)", body_txt):
    first.setdefault(f"{m.group(1)} {m.group(2)}", m.start())
fig_seq = [k for k in sorted(first, key=first.get) if k.startswith("Figure")]
tab_seq = [k for k in sorted(first, key=first.get) if k.startswith("Table")]
chk2 = dict(figure_cite_order=fig_seq, figure_insert_order=FIG_ORDER, table_cite_order=tab_seq,
            figures_ok=fig_seq == [f"Figure {i}" for i in range(1, len(fig_seq) + 1)] and FIG_ORDER == [f"Figure{i}" for i in range(1, len(FIG_ORDER) + 1)],
            tables_ok=tab_seq == [f"Table {i}" for i in range(1, len(tab_seq) + 1)],
            abstract_words=len(T.abstract(X)[1].split()), words_body=len(body_txt.split()))
print(json.dumps(chk2, indent=1))

# --- supplementary file (figures S1-S2) ---
sd = docx.Document()
st = sd.styles["Normal"]; st.font.name = "Palatino Linotype"; st.font.size = docx.shared.Pt(10)
sd.sections[0].left_margin = sd.sections[0].right_margin = docx.shared.Mm(15)
sd.add_paragraph().add_run("Supplementary Materials").bold = True
cfp = X["cf_fixed_pl"]
cap, rows, foot = T.supp_table_s1(X)
p = sd.add_paragraph(); lab = " ".join(cap.split(" ", 2)[:2]); p.add_run(lab).bold = True; p.add_run(cap[len(lab):])
tb = sd.add_table(rows=len(rows), cols=len(rows[0])); tb.style = "Table Grid"
for i, row in enumerate(rows):
    for j, v in enumerate(row):
        c = tb.cell(i, j); c.text = v
        for r_ in c.paragraphs[0].runs: r_.font.size = docx.shared.Pt(8); r_.bold = i == 0
sd.add_paragraph().add_run(foot).font.size = docx.shared.Pt(8)
for name in ["FigureS1", "FigureS2"]:
    sd.add_paragraph().add_run().add_picture(str(FIGDIR / f"{name}.png"), width=docx.shared.Mm(180))
    cap = T.SUPP_CAPTIONS[name].format(pl=cfp["cate_region"]["n_ci_below_0"], G=cfp["regions"])
    p = sd.add_paragraph(); lab = " ".join(cap.split(" ", 2)[:2]); p.add_run(lab).bold = True; p.add_run(cap[len(lab):])
sd.save(R / "manuscript/C10_Supplementary.docx")
print("supplementary written")
